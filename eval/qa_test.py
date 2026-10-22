"""Run a Q&A test set from an Excel scorecard through Radix and write the answers back.

The workbook has a sheet "Domande e valutazione" with a header row starting at "ID" and the
columns: ID, Livello, Tema, Domanda da fare al RAG, Risposta attesa, Fonte nel documento,
Pagina PDF, Risposta del RAG, Punti (0–2), Note. Each question goes through the real RAG
pipeline (rag.answer_stream, as a first chat turn); the answer and its sources fill
"Risposta del RAG", and extra columns record the cited pages, whether the expected PDF page
was cited, and the timings. Scoring (Punti) stays a human judgement.

Usage (inside the api container, like run_eval.py):
  python /app/eval/qa_test.py IN.xlsx OUT.xlsx --doc "Electric Power Transformer Engineering"
"""

import argparse
import re
import sys
import time
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from run_eval import EvalDeps, build_deps  # noqa: E402

from app.core.config import get_settings  # noqa: E402
from app.services import rag  # noqa: E402
from app.services.rag import ChatResult  # noqa: E402

SHEET = "Domande e valutazione"
EXTRA_HEADERS = ["Pagine citate", "Pagina attesa citata", "Rifiuto", "1° token (s)", "Totale (s)"]


def expected_pages(cell: Any) -> set[int]:
    """'10' → {10}; '353–354' → {353, 354}; '315; 441' → {315, 441}; '—' → set()."""
    pages: set[int] = set()
    text = re.sub(r"Cap\.?\s*[\d.]+", "", str(cell or ""))  # a chapter number is not a page
    for a, b in re.findall(r"(\d+)\s*(?:[–-]\s*(\d+))?", text):
        start = int(a)
        pages.update(range(start, int(b) + 1) if b else [start])
    return pages


def cited_pages(result: ChatResult, doc: str) -> set[int]:
    """PDF pages of `doc` the answer cites; a chunk spanning pages counts for each of them."""
    pages: set[int] = set()
    for c in result.citations:
        if c.title is None or doc.lower() not in c.title.lower():
            continue
        pages.add(c.page)
        pages.update(int(p) for p in (c.bboxes or {}) if str(p).isdigit())
    return pages


def ask(deps: EvalDeps, question: str) -> tuple[ChatResult, float | None, float]:
    start = time.perf_counter()
    first: float | None = None
    result: ChatResult | None = None
    for kind, payload in rag.answer_stream(
        db=deps.db,
        embedder=deps.embedder,
        client=deps.client,
        provider=deps.provider,
        messages=[{"role": "user", "content": question}],
        allowed_collection_ids=None,
        refusal_threshold=deps.threshold,
        query_translate=get_settings().rag_query_translate,
    ):
        if kind == "token" and first is None:
            first = time.perf_counter() - start
        elif kind == "final":
            result = payload
    assert result is not None
    return result, first, time.perf_counter() - start


def _sources(result: ChatResult) -> str:
    return "; ".join(f"[{c.n}] {c.title} p.{c.page}" for c in result.citations) or "—"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--doc", required=True, help="title (substring) of the tested document")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.input)
    ws = wb[SHEET]
    header_row = next(r for r in range(1, 20) if ws.cell(r, 1).value == "ID")
    cols = {str(ws.cell(header_row, c).value).strip(): c for c in range(1, ws.max_column + 1)}
    q_col, page_col = cols["Domanda da fare al RAG"], cols["Pagina PDF"]
    answer_col = cols["Risposta del RAG"]
    extra = ws.max_column + 1
    for i, title in enumerate(EXTRA_HEADERS):
        ws.cell(header_row, extra + i, title)

    deps = build_deps()
    rows = [r for r in range(header_row + 1, ws.max_row + 1) if ws.cell(r, q_col).value]
    hits = 0
    with_page = 0
    for n, r in enumerate(rows, 1):
        question = str(ws.cell(r, q_col).value)
        print(f"  {n}/{len(rows)} {question[:70]}", flush=True)
        result, first, total = ask(deps, question)
        want = expected_pages(ws.cell(r, page_col).value)
        got = cited_pages(result, args.doc)
        hit = bool(want & got) if want else None
        if hit is not None:
            with_page += 1
            hits += hit
        ws.cell(r, answer_col, f"{result.answer_md}\n\nFonti: {_sources(result)}")
        values = [
            ", ".join(map(str, sorted(got))) or "—",
            "—" if hit is None else ("sì" if hit else "no"),
            "sì" if result.refusal else "no",
            None if first is None else round(first, 2),
            round(total, 2),
        ]
        for i, v in enumerate(values):
            ws.cell(r, extra + i, v)

    wb.save(args.output)
    print(f"\nExpected page cited: {hits}/{with_page}. Written to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
